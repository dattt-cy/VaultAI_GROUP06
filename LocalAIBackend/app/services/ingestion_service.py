"""
Ingestion Service
=================
Luồng xử lý trọn gói từ file → vector store + SQLite:
1. Hash file để kiểm tra trùng lặp
2. Tạo Document record (PROCESSING)
3. Trích text theo định dạng (pdf/txt/docx/xlsx)
4. Chunk text → embeddings → ChromaDB
5. Cross-reference chunks vào SQLite document_pages
6. Cập nhật status SUCCESS / FAILED
"""

import json
import os
import uuid
from typing import Optional

from sqlalchemy.orm import Session

from app.crud.crud_document import create_document, create_document_page, update_document_status
from app.schemas.doc_schema import DocumentCreate, DocumentPageCreate
from app.services.document_parser import (
    generate_file_hash,
    extract_pages_from_pdf,
    chunk_text_parent_child,
    chunk_pages_parent_child,
)
from app.services.hybrid_retriever import sync_page_to_fts
from app.services.vector_store import add_documents_to_store


def _extract_text(file_path: str, file_type: str) -> str:
    """Trích text từ file theo định dạng (legacy – dùng cho non-PDF)."""
    ft = file_type.lower()

    if ft == "pdf":
        # Không dùng hàm này cho PDF nữa – dùng _extract_pages_with_metadata()
        from app.services.document_parser import extract_text_from_pdf
        return extract_text_from_pdf(file_path)

    if ft == "txt":
        for enc in ("utf-8", "utf-8-sig", "cp1258", "latin-1"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        raise ValueError("Không thể đọc file TXT: encoding không hỗ trợ")

    if ft in ("docx", "doc"):
        try:
            import docx
            doc = docx.Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            raise ValueError("Cần cài: pip install python-docx")

    if ft in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(c) if c is not None else "" for c in row)
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        except ImportError:
            raise ValueError("Cần cài: pip install openpyxl")

    raise ValueError(f"Định dạng chưa hỗ trợ: {file_type}")


def ingest_file(
    db: Session,
    file_path: str,
    filename: str,
    file_type: str,
    category_id: int,
    uploaded_by: int,
    scope: str = "PERSONAL",
    session_id: Optional[int] = None,
) -> object:
    """
    Hàm chính: nhận file đã được lưu disk, trích xuất text, chunk,
    đẩy vào ChromaDB và SQLite. Trả về Document ORM object.
    """
    # 1. Đo lường cơ bản
    file_size = os.path.getsize(file_path)
    file_hash = generate_file_hash(file_path)

    # 2. Tạo Document record với status PROCESSING
    doc_in = DocumentCreate(
        title=filename,
        file_type=file_type,
        category_id=category_id,
        file_size_bytes=file_size,
        file_hash=file_hash,
        file_path=file_path,
        uploaded_by=uploaded_by,
    )
    db_doc = create_document(db, doc_in)
    update_document_status(db, db_doc, "PROCESSING")

    # Set scope/session_id ngay lập tức để frontend thấy file trong lúc chunk
    db_doc.document_scope = scope.upper()
    if scope.upper() == "PERSONAL" and session_id:
        db_doc.session_id = session_id
    db.commit()

    try:
        # ── Phân nhánh xử lý theo định dạng ──
        ft = file_type.lower()

        if ft == "pdf":
            chunks_with_pages = _ingest_pdf(file_path, filename)
        else:
            text = _extract_text(file_path, file_type)
            if not text.strip():
                raise ValueError("File không có nội dung văn bản để xử lý")
            chunks_with_pages = chunk_text_parent_child(text)

        if not chunks_with_pages:
            raise ValueError("File không có nội dung văn bản để xử lý")

        # ── Lưu parent-child vào SQLite + ChromaDB ──
        documents_for_vector: list = []
        metadatas_for_vector: list = []
        chroma_ids_input: list = []

        total_tokens = _save_parent_child_chunks(
            db=db,
            chunks=chunks_with_pages,
            db_doc=db_doc,
            filename=filename,
            documents_for_vector=documents_for_vector,
            metadatas_for_vector=metadatas_for_vector,
            chroma_ids_input=chroma_ids_input,
        )

        if documents_for_vector:
            add_documents_to_store(
                texts=documents_for_vector,
                metadatas=metadatas_for_vector,
                ids=chroma_ids_input,
            )

        n_parents = sum(1 for c in chunks_with_pages if c["chunk_type"] == "parent")
        n_children = sum(1 for c in chunks_with_pages if c["chunk_type"] == "child")
        update_document_status(db, db_doc, "SUCCESS", total_tokens=total_tokens)
        print(f"[Ingestion OK] {filename} → {n_parents} parents, {n_children} children, {total_tokens} tokens")

    except Exception as e:
        update_document_status(db, db_doc, "FAILED", error=str(e))
        print(f"[Ingestion FAIL] {filename}: {e}")
        raise

    return db_doc


def _ingest_pdf(file_path: str, filename: str) -> list[dict]:
    """PDF: trích text từng trang → parent-child chunking, giữ page_number thật."""
    pages = extract_pages_from_pdf(file_path)
    if not pages:
        raise ValueError("PDF không có nội dung văn bản để xử lý")
    chunks = chunk_pages_parent_child(pages)
    n_parents = sum(1 for c in chunks if c["chunk_type"] == "parent")
    print(f"[PDF Parse] {filename}: {len(pages)} trang → {n_parents} parents")
    return chunks


def _save_parent_child_chunks(
    db: Session,
    chunks: list[dict],
    db_doc,
    filename: str,
    documents_for_vector: list,
    metadatas_for_vector: list,
    chroma_ids_input: list,
) -> int:
    """
    Lưu parent-child chunks vào SQLite.
    Pass 1: lưu parent rows → lấy DB ids.
    Pass 2: lưu child rows với parent_chunk_id → thêm vào ChromaDB batch.
    Chỉ child rows được insert vào ChromaDB và FTS5.
    Trả về tổng token count (tính từ child chunks).
    """
    parent_rows = [c for c in chunks if c["chunk_type"] == "parent"]
    child_rows  = [c for c in chunks if c["chunk_type"] == "child"]

    # Pass 1 — lưu parent rows
    parent_index_to_db_id: dict[int, int] = {}
    global_chunk_index = 0

    for p in parent_rows:
        v_id = str(uuid.uuid4())
        child_count = sum(1 for c in child_rows if c["parent_index"] == p["parent_index"])
        page_in = DocumentPageCreate(
            document_id=db_doc.id,
            chunk_index=global_chunk_index,
            raw_content=p["text"],
            token_count=len(p["text"]) // 4,
            page_metadata=json.dumps({
                "source": filename,
                "document_id": db_doc.id,
                "chunk_index": global_chunk_index,
                "page_number": p["page_number"],
                "vector_id": v_id,
            }),
            vector_id=v_id,
            chunk_type="parent",
            parent_chunk_id=None,
            child_count=child_count,
        )
        db_page = create_document_page(db, page_in)
        parent_index_to_db_id[p["parent_index"]] = db_page.id
        global_chunk_index += 1

    # Pass 2 — lưu child rows, build ChromaDB batch
    total_tokens = 0
    for c in child_rows:
        parent_db_id = parent_index_to_db_id.get(c["parent_index"])
        v_id = str(uuid.uuid4())
        tokens = len(c["text"]) // 4
        total_tokens += tokens
        meta = {
            "source": filename,
            "document_id": db_doc.id,
            "chunk_index": global_chunk_index,
            "page_number": c["page_number"],
            "vector_id": v_id,
            "parent_chunk_db_id": parent_db_id,
        }
        page_in = DocumentPageCreate(
            document_id=db_doc.id,
            chunk_index=global_chunk_index,
            raw_content=c["text"],
            token_count=tokens,
            page_metadata=json.dumps(meta),
            vector_id=v_id,
            chunk_type="child",
            parent_chunk_id=parent_db_id,
        )
        db_page = create_document_page(db, page_in)
        sync_page_to_fts(db, db_page.id, c["text"])

        documents_for_vector.append(c["text"])
        metadatas_for_vector.append(meta)
        chroma_ids_input.append(v_id)
        global_chunk_index += 1

    return total_tokens
